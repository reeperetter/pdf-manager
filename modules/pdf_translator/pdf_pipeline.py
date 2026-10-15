"""
Основний PDF-пайплайн: рендеринг сторінок, вставка розпізнаного/
перекладеного тексту в PDF, конвертація зображень у PDF, збереження
окремого текстового файлу (DOCX/TXT).

Сам рушій розпізнавання (Tesseract) та розпрямлення сторінок - у ocr.py.
Переклад виконується онлайн через незалежні сервіси (MyMemory / DeepL).
"""
import gc
import io
import os
import time

import docx as _docx_module
import fitz  # PyMuPDF
from deep_translator import DeeplTranslator, MyMemoryTranslator
from PIL import Image, ImageOps

from .ocr import (detect_and_fix_orientation, dewarp_page_image,
                 ocr_lines_from_image)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
# DejaVuSans.ttf фізично лежить у КОРЕНІ репозиторію (поруч із main.py), а
# не в modules/pdf_translator/. У зібраному --onedir/--onefile застосунку
# це не впадало в очі, бо build.yml явно копіює файл саме сюди через
# --add-data "DejaVuSans.ttf;modules/pdf_translator" - але при запуску з
# сирого коду такого копіювання ніхто не робить, SCRIPT_DIR/DejaVuSans.ttf
# просто не існує, і пошук одразу провалювався в FALLBACK_FONTS (жорстко
# прописані системні шляхи, яких на конкретній машині може й не бути -
# звідси "Не знайдено TTF-шрифт з підтримкою кирилиці" при запуску з коду).
# Перевіряємо обидва місця, щоб працювало і так, і так, незалежно від
# того, чи є той самий --add-data у build.yml.
_PROJECT_ROOT = os.path.dirname(os.path.dirname(SCRIPT_DIR))
_BUNDLED_FONT_CANDIDATES = [
    os.path.join(SCRIPT_DIR, "DejaVuSans.ttf"),
    os.path.join(_PROJECT_ROOT, "DejaVuSans.ttf"),
]
BUNDLED_FONT = next((p for p in _BUNDLED_FONT_CANDIDATES if os.path.isfile(p)), None)

FALLBACK_FONTS = [
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\calibri.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",         # шлях на Fedora/Nobara
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/System/Library/Fonts/Helvetica.ttc",
]


class ProcessingCancelled(Exception):
    pass


def interruptible_sleep(duration, cancel_event=None, step=0.05):
    """Виконує паузу з можливістю миттєвого переривання через cancel_event."""
    end_time = time.time() + duration
    while time.time() < end_time:
        if cancel_event and cancel_event.is_set():
            raise ProcessingCancelled()
        time.sleep(min(step, max(0, end_time - time.time())))


def find_unicode_font(user_path=None):
    """Повертає шлях до TTF-шрифту, що підтримує кирилицю."""
    if user_path and os.path.isfile(user_path):
        return user_path
    if BUNDLED_FONT and os.path.isfile(BUNDLED_FONT):
        return BUNDLED_FONT
    for p in FALLBACK_FONTS:
        if os.path.isfile(p):
            return p
    return None


def extract_lines_from_page(page):
    """Витягує рядки тексту з наявного текстового шару PDF."""
    lines = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            if not spans:
                continue
            text = "".join(s["text"] for s in spans).strip()
            if not text:
                continue
            x0 = min(s["bbox"][0] for s in spans)
            y0 = min(s["bbox"][1] for s in spans)
            x1 = max(s["bbox"][2] for s in spans)
            y1 = max(s["bbox"][3] for s in spans)
            lines.append(
                {"text": text, "bbox": (x0, y0, x1, y1), "page_coords": True}
            )
    return lines


def line_to_rect(line, scale_x, scale_y):
    """Перетворює bbox рядка у fitz.Rect."""
    x0, y0, x1, y1 = line["bbox"]
    if line.get("page_coords"):
        return fitz.Rect(x0, y0, x1, y1)
    return fitz.Rect(x0 * scale_x, y0 * scale_y, x1 * scale_x, y1 * scale_y)


def line_to_pixel_bbox(line, scale_x, scale_y):
    """Перетворює bbox рядка у координати пікселів зображення."""
    x0, y0, x1, y1 = line["bbox"]
    if line.get("page_coords"):
        return (x0 / scale_x, y0 / scale_y, x1 / scale_y, y1 / scale_y)
    return line["bbox"]


def fit_fontsize(text, rect_width, rect_height, fontfile, max_size=None):
    """Підбирає розмір шрифту під прямокутник."""
    if max_size is None:
        max_size = max(4, rect_height * 0.85)
    size = max_size
    try:
        font = fitz.Font(fontfile=fontfile)
    except Exception:
        font = fitz.Font("helv")
    while size > 4:
        width = font.text_length(text, fontsize=size)
        if width <= rect_width or size <= 4.5:
            break
        size -= 0.5
    return size


def sample_bg_color(pil_img, bbox, pad=4):
    """Визначає колір фону навколо текстового блоку."""
    x0, y0, x1, y1 = bbox
    w, h = pil_img.size
    x0i, y0i, x1i, y1i = int(x0), int(y0), int(x1), int(y1)

    strips = [
        (x0i, max(0, y0i - pad), x1i, y0i),
        (x0i, y1i, x1i, min(h, y1i + pad)),
        (max(0, x0i - pad), y0i, x0i, y1i),
        (x1i, y0i, min(w, x1i + pad), y1i),
    ]

    samples_r, samples_g, samples_b = [], [], []
    for sx, sy, ex, ey in strips:
        if ex <= sx or ey <= sy:
            continue
        crop = pil_img.crop((sx, sy, ex, ey)).convert("RGB")
        for p in crop.getdata():
            samples_r.append(p[0])
            samples_g.append(p[1])
            samples_b.append(p[2])

    if not samples_r:
        return (1, 1, 1)

    def median(vals):
        vals = sorted(vals)
        n = len(vals)
        mid = n // 2
        if n % 2:
            return vals[mid]
        return (vals[mid - 1] + vals[mid]) / 2

    return (median(samples_r) / 255, median(samples_g) / 255, median(samples_b) / 255)


IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")


def _flatten_image_for_pdf(img):
    """Готує зображення для вставки у PDF-сторінку.

    Повертає (img_rgb_or_rgba, has_real_transparency):
    - якщо в зображення є альфа-канал, і хоч один піксель дійсно не
      повністю непрозорий (alpha < 255 десь є) - повертає RGBA як є,
      True. Такі зображення далі зберігаються як PNG (JPEG альфу не
      підтримує), і прозорі ділянки лишаються прозорими на сторінці
      (переглядач покаже під ними фон самої сторінки PDF, типово
      білий) - це і є "прозорий, де можливо".
    - інакше - складає зображення на БІЛИЙ фон і повертає RGB, False.
      Раніше тут був просто img.convert("RGB") - для RGBA-зображень
      PIL при такій конвертації просто ВІДКИДАЄ альфа-канал і лишає
      "сирі" RGB-значення пікселів як є, а в стертих/прозорих ділянках
      типового PNG ці значення часто (0, 0, 0) - звідси й чорний фон
      замість очікуваного білого чи прозорого.
    """
    has_alpha = img.mode in ("RGBA", "LA") or (
        img.mode == "P" and "transparency" in img.info
    )
    if not has_alpha:
        return img.convert("RGB"), False

    img = img.convert("RGBA")
    alpha_extrema = img.getchannel("A").getextrema()
    is_really_transparent = alpha_extrema[0] < 255
    if is_really_transparent:
        return img, True

    # Альфа-канал технічно є, але всі пікселі повністю непрозорі -
    # звичайний RGB-шлях, без зайвого PNG.
    background = Image.new("RGB", img.size, (255, 255, 255))
    background.paste(img, mask=img.getchannel("A"))
    return background, False


def images_to_pdf(image_paths, output_path, log_fn=None, progress_fn=None,
                  cancel_event=None, jpeg_quality=88, max_dimension=2200):
    """Збирає список зображень в один PDF."""
    assumed_dpi = 200
    doc = fitz.open()
    total_before = 0
    total_after = 0
    try:
        for i, img_path in enumerate(image_paths):
            if cancel_event is not None and cancel_event.is_set():
                raise ProcessingCancelled()
            try:
                total_before += os.path.getsize(img_path)
            except OSError:
                pass

            img = Image.open(img_path)
            img = ImageOps.exif_transpose(img)
            img, is_transparent = _flatten_image_for_pdf(img)

            if max(img.width, img.height) > max_dimension:
                scale = max_dimension / max(img.width, img.height)
                img = img.resize(
                    (max(1, int(img.width * scale)),
                     max(1, int(img.height * scale))),
                    Image.LANCZOS,
                )

            w_pt = img.width * 72 / assumed_dpi
            h_pt = img.height * 72 / assumed_dpi
            page = doc.new_page(width=w_pt, height=h_pt)

            buf = io.BytesIO()
            if is_transparent:
                # JPEG не підтримує прозорість - для таких зображень
                # зберігаємо без втрат у PNG, щоб прозорі ділянки
                # лишились прозорими на сторінці.
                img.save(buf, format="PNG", optimize=True)
            else:
                img.save(buf, format="JPEG", quality=jpeg_quality, optimize=True)
            total_after += buf.tell()
            page.insert_image(fitz.Rect(0, 0, w_pt, h_pt),
                              stream=buf.getvalue())

            if log_fn:
                log_fn(
                    f"  Додано сторінку {i + 1}/{len(image_paths)}: "
                    f"{os.path.basename(img_path)}"
                )
            if progress_fn:
                progress_fn(i + 1, len(image_paths))
        doc.save(output_path, garbage=3, deflate=True)
    finally:
        doc.close()

    if log_fn and total_before > 0:
        final_size = os.path.getsize(output_path)
        log_fn(
            f"  Стиснення: вихідні зображення ~{total_before / 1_048_576:.1f} МБ "
            f"→ підсумковий PDF {final_size / 1_048_576:.1f} МБ"
        )
    return output_path


def detect_text_language(text):
    """Спроба визначити мову тексту локально без зовнішніх API."""
    if not text or len(text.strip()) < 3:
        return "en"

    try:
        from langdetect import detect
        lang = detect(text)
        return lang if lang else "en"
    except Exception:
        pass

    # Резервний аналіз за алфавітом
    has_cyrillic = any('\u0400' <= char <= '\u04FF' for char in text)
    if has_cyrillic:
        return "uk"
    return "en"


def translate_batch_microsoft(lines_list, target_lang="uk", source_lang="auto", log_fn=None, cancel_event=None, deepl_api_key=None):
    """
    Онлайн-переклад через MyMemory або DeepL.
    Автоматично замінює 'auto' на визначений ISO-код перед відправкою у MyMemory.
    """
    if not lines_list:
        return []

    if cancel_event and cancel_event.is_set():
        raise ProcessingCancelled()

    # 1. Якщо передано DeepL — він підтримує 'auto'
    if deepl_api_key:
        try:
            translator = DeeplTranslator(api_key=deepl_api_key, source=source_lang, target=target_lang, use_free_api=True)
            return [translator.translate(line) if line.strip() else "" for line in lines_list]
        except Exception as e:
            if log_fn:
                log_fn(f"  [!] DeepL помилка: {e}. Перехід на MyMemory...")

    # 2. Якщо source_lang == 'auto' або це комбінація кількох мов Tesseract
    #    (наприклад "ukr+eng"), однозначного коду немає - визначаємо мову
    #    локально за текстом.
    actual_source = source_lang
    if actual_source == "auto" or (actual_source and "+" in actual_source):
        sample_text = " ".join([l.strip() for l in lines_list if l.strip()][:5])
        actual_source = detect_text_language(sample_text)
        if log_fn:
            log_fn(f"  [i] Автовизначення мови джерела: '{actual_source}'")

    # Мапінг кодів під вимоги MyMemory (потрібні повні локалі, напр. "uk-UA").
    # Комбобокс вибору мови в gui.py віддає 3-літерні коди Tesseract
    # ("eng", "ukr", ...), а detect_text_language/langdetect - короткі
    # ISO-коди ("en", "uk", ...). Мапа повинна покривати ОБИДВА формати,
    # інакше MyMemory відповідає "No support for the provided language"
    # і переклад мовчки пропускається (файл зберігається неперекладеним).
    lang_map = {
        "uk": "uk-UA", "ukr": "uk-UA",
        "en": "en-US", "eng": "en-US",
        "de": "de-DE", "deu": "de-DE", "ger": "de-DE",
        "fr": "fr-FR", "fra": "fr-FR", "fre": "fr-FR",
        "es": "es-ES", "spa": "es-ES",
        "it": "it-IT", "ita": "it-IT",
        "pl": "pl-PL", "pol": "pl-PL",
        "ru": "ru-RU", "rus": "ru-RU",
    }

    tgt = lang_map.get(target_lang, target_lang)
    src = lang_map.get(actual_source, actual_source)

    if log_fn:
        log_fn(f"  [i] MyMemory: source='{src}' target='{tgt}'")

    try:
        translator = MyMemoryTranslator(source=src, target=tgt)
        translated_lines = []

        for line in lines_list:
            if cancel_event and cancel_event.is_set():
                raise ProcessingCancelled()

            text = line.strip()
            if not text:
                translated_lines.append("")
                continue

            interruptible_sleep(0.1, cancel_event=cancel_event)

            try:
                res = translator.translate(text)
                if res and "INVALID SOURCE LANGUAGE" in str(res).upper():
                    translated_lines.append(text)
                else:
                    translated_lines.append(res if res else text)
            except Exception as item_err:
                if log_fn:
                    log_fn(f"  [!] Помилка рядка: {item_err}")
                translated_lines.append(text)

        return translated_lines

    except Exception as e:
        if log_fn:
            log_fn(f"  [!] MyMemory Помилка: {e}")

    return lines_list


def _write_txt_file(path, page_texts):
    with open(path, "w", encoding="utf-8") as f:
        for i, text in enumerate(page_texts):
            if len(page_texts) > 1:
                if i > 0:
                    f.write("\n\n")
                f.write(f"----- Сторінка {i + 1} -----\n\n")
            f.write(text)
            f.write("\n")


def _write_docx_file(path, page_texts):
    doc = _docx_module.Document()
    for i, text in enumerate(page_texts):
        if i > 0:
            doc.add_page_break()
        if len(page_texts) > 1:
            heading = doc.add_paragraph(f"Сторінка {i + 1}")
            heading.runs[0].bold = True
        for line in text.split("\n"):
            if line.strip():
                doc.add_paragraph(line)
    doc.save(path)


def save_text_outputs(page_texts, output_dir, base_name, suffix, formats, split_pages):
    saved = []
    writers = {"txt": _write_txt_file, "docx": _write_docx_file}
    if split_pages:
        for i, text in enumerate(page_texts):
            for fmt in formats:
                out_path = os.path.join(
                    output_dir, f"{base_name}{suffix}_p{i + 1}.{fmt}")
                writers[fmt](out_path, [text])
                saved.append(out_path)
    else:
        for fmt in formats:
            out_path = os.path.join(output_dir, f"{base_name}{suffix}.{fmt}")
            writers[fmt](out_path, page_texts)
            saved.append(out_path)
    return saved


class PdfProcessResult:
    """Результат обробки одного PDF у пам'яті (ще не збережений на диск).

    Документи (searchable / translated[lang]) належить викликачу: його
    справа або зберегти їх, або вставити в інший документ і закрити."""

    def __init__(self, searchable, translated, original_texts,
                 translated_texts, low_confidence_pages):
        self.searchable = searchable
        self.translated = translated
        self.original_texts = original_texts
        self.translated_texts = translated_texts
        self.low_confidence_pages = low_confidence_pages

    def close(self):
        if self.searchable is not None:
            self.searchable.close()
            self.searchable = None
        for d in self.translated.values():
            d.close()
        self.translated = {}


def _process_pdf_core(
    input_path,
    dpi,
    make_searchable,
    translate_targets,
    font_path,
    log_fn,
    cancel_event,
    page_start=None,
    page_end=None,
    page_progress_fn=None,
    glossary=None,
    low_confidence_threshold=55,
    dewarp=False,
    collect_text=False,
    source_lang="auto",
    deepl_api_key=None,
):
    """Розпізнає/перекладає один PDF і повертає PdfProcessResult (у пам'яті)."""
    export_text_formats = {"txt"} if collect_text else set()
    glossary = glossary or set()
    low_confidence_pages = []
    original_page_texts = []
    translated_page_texts = {lang: [] for lang in translate_targets}

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    doc = fitz.open(input_path)
    n_pages = doc.page_count

    fontfile = find_unicode_font(font_path)
    if fontfile is None:
        doc.close()
        raise RuntimeError("Не знайдено TTF-шрифт з підтримкою кирилиці.")

    out_searchable = fitz.open() if make_searchable else None
    out_translated = {lang: fitz.open() for lang in translate_targets}

    try:
        for page_index in range(n_pages):
            if cancel_event and cancel_event.is_set():
                raise ProcessingCancelled()

            page_num = page_index + 1
            in_range = (page_start is None or page_num >= page_start) and \
                       (page_end is None or page_num <= page_end)

            log_fn(f"[{base_name}] Сторінка {page_num}/{n_pages}: рендеринг...")
            page = doc[page_index]
            pix = page.get_pixmap(dpi=dpi)
            img_bytes = pix.tobytes("png")
            pil_img = Image.open(io.BytesIO(img_bytes))

            was_rotated = False
            if in_range:
                pil_img, was_rotated = detect_and_fix_orientation(
                    pil_img, log_fn=log_fn)
                if was_rotated:
                    buf = io.BytesIO()
                    pil_img.save(buf, format="PNG")
                    img_bytes = buf.getvalue()

            page_w = pil_img.width * 72 / dpi
            page_h = pil_img.height * 72 / dpi
            scale_x = page_w / pil_img.width
            scale_y = page_h / pil_img.height

            if dewarp and in_range:
                try:
                    corrected = dewarp_page_image(
                        pil_img, log_fn=log_fn, cancel_event=cancel_event)
                    if corrected is not pil_img:
                        pil_img = corrected
                        buf = io.BytesIO()
                        pil_img.save(buf, format="PNG")
                        img_bytes = buf.getvalue()
                except RuntimeError as e:
                    log_fn(f"  [!] Розпрямлення недоступне: {e}")
                    dewarp = False

            lines = []
            if in_range and (make_searchable or translate_targets or collect_text):
                if was_rotated:
                    log_fn(
                        f"[{base_name}] Сторінка {page_num}/{n_pages}: "
                        f"орієнтацію змінено — розпізнавання (OCR) за новими координатами...")
                    lines, ocr_score = ocr_lines_from_image(pil_img, source_lang=source_lang)
                    if lines and ocr_score < low_confidence_threshold:
                        low_confidence_pages.append(page_num)
                        log_fn(
                            f"  [!] Сторінка {page_num}: невисока впевненість "
                            f"OCR ({ocr_score:.0f}/100) - варто перевірити вручну")
                else:
                    lines = extract_lines_from_page(page)
                    if lines:
                        log_fn(
                            f"[{base_name}] Сторінка {page_num}/{n_pages}: "
                            f"текст з PDF ({len(lines)} рядків), OCR пропущено")
                    else:
                        log_fn(
                            f"[{base_name}] Сторінка {page_num}/{n_pages}: "
                            f"немає текстового шару — розпізнавання (OCR)...")
                        lines, ocr_score = ocr_lines_from_image(pil_img, source_lang=source_lang)
                        if lines and ocr_score < low_confidence_threshold:
                            low_confidence_pages.append(page_num)
                            log_fn(
                                f"  [!] Сторінка {page_num}: невисока впевненість "
                                f"OCR ({ocr_score:.0f}/100) - варто перевірити вручну")
            elif not in_range:
                log_fn(
                    f"[{base_name}] Сторінка {page_num}/{n_pages}: поза вибраним "
                    f"діапазоном, пропущено (сторінку скопійовано без обробки)")

            if export_text_formats:
                original_page_texts.append("\n".join(ln["text"] for ln in lines))

            # 1) Searchable PDF
            if out_searchable is not None:
                sp = out_searchable.new_page(width=page_w, height=page_h)
                sp.insert_image(sp.rect, stream=img_bytes)
                for line in lines:
                    rect = line_to_rect(line, scale_x, scale_y)
                    if rect.is_empty or rect.width <= 0 or rect.height <= 0:
                        continue
                    fs = fit_fontsize(
                        line["text"], rect.width, rect.height, fontfile)
                    baseline = fitz.Point(rect.x0, rect.y1 - 0.15 * fs)
                    try:
                        sp.insert_text(
                            baseline,
                            line["text"],
                            fontsize=fs,
                            fontfile=fontfile,
                            fontname="cyr1",
                            render_mode=3,
                        )
                    except Exception as e:
                        log_fn(f"  [!] Пропущено рядок (текстовий шар): {e}")

            # 2) Translated PDF
            if translate_targets:
                translations = {}
                if lines:
                    log_fn(f"[{base_name}] Сторінка {page_num}/{n_pages}: переклад ({len(lines)} рядків)...")
                    raw_texts = [ln["text"] for ln in lines]

                    for lang in translate_targets:
                        translated_lines = translate_batch_microsoft(
                            raw_texts, target_lang=lang, source_lang=source_lang,
                            log_fn=log_fn, cancel_event=cancel_event, deepl_api_key=deepl_api_key
                        )
                        translations[lang] = translated_lines

                if export_text_formats:
                    for lang in translate_targets:
                        text = "\n".join(translations.get(
                            lang, [])) if lines else ""
                        translated_page_texts[lang].append(text)

                for lang in translate_targets:
                    tp = out_translated[lang].new_page(width=page_w, height=page_h)
                    tp.insert_image(tp.rect, stream=img_bytes)
                    if not lines:
                        continue
                    translated_lines = translations[lang]

                    to_draw = []
                    for line, tr_text in zip(lines, translated_lines):
                        orig_text = line["text"].strip()
                        tr_text_stripped = (tr_text or "").strip()
                        if tr_text_stripped.lower() == orig_text.lower():
                            continue
                        rect = line_to_rect(line, scale_x, scale_y)
                        if rect.is_empty or rect.width <= 0 or rect.height <= 0:
                            continue
                        px_bbox = line_to_pixel_bbox(line, scale_x, scale_y)
                        bg = sample_bg_color(pil_img, px_bbox)
                        cover = fitz.Rect(rect.x0 - 1, rect.y0 - 1,
                                          rect.x1 + 1, rect.y1 + 1)
                        tp.draw_rect(cover, color=bg, fill=bg, overlay=True)
                        to_draw.append((rect, tr_text))

                    for rect, tr_text in to_draw:
                        if not tr_text:
                            continue
                        fs = fit_fontsize(tr_text, rect.width,
                                          rect.height, fontfile)
                        baseline = fitz.Point(rect.x0, rect.y1 - 0.15 * fs)
                        try:
                            tp.insert_text(
                                baseline,
                                tr_text,
                                fontsize=fs,
                                fontfile=fontfile,
                                fontname="cyr1",
                                color=(0, 0, 0),
                            )
                        except Exception as e:
                            log_fn(f"  [!] Пропущено рядок (переклад {lang}): {e}")

            del pix, pil_img, img_bytes, lines
            page = None
            fitz.TOOLS.store_shrink(100)
            gc.collect()

            if page_progress_fn:
                page_progress_fn(page_num, n_pages)

    except BaseException:
        # Скасування чи помилка посеред файлу - не лишаємо відкритих документів.
        if out_searchable is not None:
            out_searchable.close()
        for d in out_translated.values():
            d.close()
        doc.close()
        raise

    doc.close()
    gc.collect()
    return PdfProcessResult(
        out_searchable, out_translated, original_page_texts,
        translated_page_texts, low_confidence_pages)


def process_pdf(
    input_path,
    output_dir,
    dpi,
    make_searchable,
    translate_targets,
    font_path,
    log_fn,
    cancel_event,
    export_text_formats=None,
    export_text_split_pages=False,
    export_original_text=False,
    **core_kwargs,
):
    """Обробляє один PDF і зберігає результати у output_dir
    (<ім'я>_searchable.pdf, <ім'я>_<мова>.pdf, текстові файли)."""
    export_text_formats = export_text_formats or set()
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    res = _process_pdf_core(
        input_path, dpi, make_searchable, translate_targets, font_path,
        log_fn, cancel_event, collect_text=bool(export_text_formats),
        **core_kwargs)
    try:
        os.makedirs(output_dir, exist_ok=True)
        saved_files = []

        if res.searchable is not None:
            out_path = os.path.join(output_dir, f"{base_name}_searchable.pdf")
            res.searchable.save(out_path, garbage=3, deflate=True)
            saved_files.append(out_path)

        for lang, d in res.translated.items():
            out_path = os.path.join(output_dir, f"{base_name}_{lang}.pdf")
            d.save(out_path, garbage=3, deflate=True)
            saved_files.append(out_path)

        if export_text_formats:
            if export_original_text:
                saved_files += save_text_outputs(
                    res.original_texts, output_dir, base_name, "_text",
                    export_text_formats, export_text_split_pages)
            for lang in translate_targets:
                saved_files += save_text_outputs(
                    res.translated_texts[lang], output_dir, base_name,
                    f"_{lang}_text", export_text_formats,
                    export_text_split_pages)
        low_conf = res.low_confidence_pages
    finally:
        res.close()
    gc.collect()
    return saved_files, low_conf


class CombinedResult:
    """Накопичує результати кількох вхідних PDF в ОДИН вихідний файл
    (на кожну мову перекладу - свій), щоб зберегти їх разом."""

    def __init__(self, make_searchable, translate_targets, export_original_text=False):
        self.export_original_text = export_original_text
        self.translate_targets = list(translate_targets)
        self.searchable = fitz.open() if make_searchable else None
        self.translated = {lang: fitz.open() for lang in self.translate_targets}
        self.original_texts = []
        self.translated_texts = {lang: [] for lang in self.translate_targets}
        self.files_added = 0

    def add(self, res):
        """Дописує результат одного файлу в кінець і закриває його."""
        try:
            if self.searchable is not None and res.searchable is not None:
                self.searchable.insert_pdf(res.searchable)
            for lang in self.translate_targets:
                if lang in res.translated:
                    self.translated[lang].insert_pdf(res.translated[lang])
                self.translated_texts[lang].extend(
                    res.translated_texts.get(lang, []))
            self.original_texts.extend(res.original_texts)
            self.files_added += 1
        finally:
            res.close()

    def save(self, out_path, export_text_formats=None, export_text_split_pages=False):
        """Зберігає накопичене. out_path - шлях до головного .pdf.

        Для розпізнавання це і є вихідний файл. Для перекладу однією
        мовою - теж точно out_path; для кількох мов до імені додається
        _<мова> (<ім'я>_uk.pdf, <ім'я>_en.pdf)."""
        export_text_formats = export_text_formats or set()
        out_dir = os.path.dirname(out_path) or "."
        stem = os.path.splitext(os.path.basename(out_path))[0]
        os.makedirs(out_dir, exist_ok=True)
        saved = []

        if self.searchable is not None:
            self.searchable.save(out_path, garbage=3, deflate=True)
            saved.append(out_path)

        for lang in self.translate_targets:
            if len(self.translate_targets) == 1 and self.searchable is None:
                lang_path = out_path
            else:
                lang_path = os.path.join(out_dir, f"{stem}_{lang}.pdf")
            self.translated[lang].save(lang_path, garbage=3, deflate=True)
            saved.append(lang_path)

        if export_text_formats:
            if self.export_original_text:
                saved += save_text_outputs(
                    self.original_texts, out_dir, stem, "_text",
                    export_text_formats, export_text_split_pages)
            for lang in self.translate_targets:
                saved += save_text_outputs(
                    self.translated_texts[lang], out_dir, stem,
                    f"_{lang}_text", export_text_formats,
                    export_text_split_pages)
        return saved

    def close(self):
        if self.searchable is not None:
            self.searchable.close()
            self.searchable = None
        for d in self.translated.values():
            d.close()
        self.translated = {}
